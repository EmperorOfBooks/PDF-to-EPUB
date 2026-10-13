from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document
from PIL import Image

from cleaner import FlowImageUnit, FlowTableUnit, FlowTextUnit
from office_extractor import OfficeExtractionError, OfficeExtractor
from router import DocumentRouter


def make_docx(path):
    document = Document()
    document.add_heading("Chapter 1", level=1)
    document.add_paragraph("Office body text.")
    image_buffer = BytesIO()
    Image.new("RGB", (32, 24), "red").save(image_buffer, format="PNG")
    document.add_picture(BytesIO(image_buffer.getvalue()))
    document.add_heading("A subsection", level=2)
    document.add_paragraph("More body text.")
    document.save(path)


def test_docx_extracts_semantic_blocks_and_embedded_images(tmp_path):
    source = tmp_path / "sample.docx"
    make_docx(source)

    cleaned = OfficeExtractor(source).extract()
    items = [item for chapter in cleaned.chapters for item in chapter.items]

    assert any(isinstance(item, FlowTextUnit) and item.kind == "heading" and item.is_chapter_heading for item in items)
    assert any(isinstance(item, FlowTextUnit) and item.kind == "paragraph" and "Office body" in item.text for item in items)
    assert any(isinstance(item, FlowImageUnit) and item.full_width for item in items)


def test_router_dispatches_docx_to_office_extractor(tmp_path):
    source = tmp_path / "sample.docx"
    make_docx(source)

    cleaned = DocumentRouter().extract(source)

    assert cleaned.chapters
    assert cleaned.chapters[0].items


def test_docx_extracts_table_in_document_order(tmp_path):
    source = tmp_path / "with_table.docx"
    document = Document()
    document.add_heading("Chapter 1", level=1)
    document.add_paragraph("Intro text.")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Header A"
    table.rows[0].cells[1].text = "Header B"
    table.rows[1].cells[0].text = "Val 1"
    table.rows[1].cells[1].text = "Val 2"
    document.add_paragraph("After table.")
    document.save(source)

    cleaned = OfficeExtractor(source).extract()
    items = [item for chapter in cleaned.chapters for item in chapter.items]
    tables = [item for item in items if isinstance(item, FlowTableUnit)]

    assert len(tables) == 1
    assert tables[0].rows == (("Header A", "Header B"), ("Val 1", "Val 2"))
    assert tables[0].has_header_row


def test_doc_extraction_without_libreoffice_raises_clear_error(tmp_path, monkeypatch):
    import office_extractor

    monkeypatch.setattr(office_extractor.shutil, "which", lambda name: None)
    source = tmp_path / "legacy.doc"
    source.write_bytes(b"not a real doc file")

    with pytest.raises(OfficeExtractionError, match="LibreOffice"):
        OfficeExtractor(source).extract()


def test_zip_slip_docx_payload_raises_security_exception(tmp_path):
    import zipfile

    evil_docx = tmp_path / "malicious.docx"
    with zipfile.ZipFile(evil_docx, "w") as z:
        z.writestr("../../evil.sh", b"echo hacked")
        z.writestr("word/document.xml", b"<xml></xml>")

    with pytest.raises(OfficeExtractionError, match="[Ss]ecurity|[Zz]ip [Ss]lip|traversal"):
        OfficeExtractor(evil_docx).extract()


def test_decompression_bomb_docx_raises_security_exception(tmp_path):
    import zipfile

    bomb_docx = tmp_path / "bomb.docx"
    with zipfile.ZipFile(bomb_docx, "w", compression=zipfile.ZIP_DEFLATED) as z:
        # Create an entry exceeding 50 MB
        z.writestr("word/document.xml", b"0" * (52 * 1024 * 1024))

    with pytest.raises(OfficeExtractionError, match="[Ss]ecurity|[Dd]ecompression bomb|[Ee]xceeded"):
        OfficeExtractor(bomb_docx).extract()


def test_command_injection_in_filename_sanitized(tmp_path, monkeypatch):
    import subprocess
    from pathlib import Path
    import office_extractor

    monkeypatch.setattr(office_extractor.shutil, "which", lambda name: "soffice")

    called_cmds = []

    def mock_run(cmd, **kwargs):
        called_cmds.append((cmd, kwargs))
        # Create a mock converted output docx file matching the temp file stem
        temp_input = Path(cmd[-1])
        converted = temp_input.parent / f"{temp_input.stem}.docx"
        make_docx(converted)
        return subprocess.CompletedProcess(cmd, returncode=0, stdout="", stderr="")

    monkeypatch.setattr(office_extractor.subprocess, "run", mock_run)

    injection_filename = "doc_name;calc.exe;&rm -rf;.doc"
    source = tmp_path / injection_filename
    source.write_bytes(b"fake doc content")

    cleaned = OfficeExtractor(source).extract()
    assert cleaned.chapters

    # Assert subprocess was called with shell=False and sanitized uuid filename
    assert len(called_cmds) == 1
    cmd, kwargs = called_cmds[0]
    assert kwargs.get("shell") is False
    assert kwargs.get("timeout") == 120
    # Command argument must not contain the injection string
    input_arg = cmd[-1]
    assert injection_filename not in input_arg
    assert "calc.exe" not in input_arg

