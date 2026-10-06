from __future__ import annotations

import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any, Iterable

import ftfy
import pypandoc
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from lxml import html as lxml_html

from cleaner import ChapterContent, CleanedDocument, FlowImageUnit, FlowTableUnit, FlowTextUnit


class OfficeExtractionError(RuntimeError):
    pass


class OfficeExtractor:
    """Convert OOXML and Pandoc-supported office formats into the shared flow model."""

    def __init__(self, input_path: str | Path) -> None:
        self.input_path = Path(input_path).expanduser().resolve()

    def extract(self) -> CleanedDocument:
        suffix = self.input_path.suffix.lower()
        if suffix == ".docx":
            return self._extract_docx()
        if suffix == ".doc":
            return self._extract_doc_via_libreoffice()
        if suffix in {".odt", ".rtf"}:
            return self._extract_with_pandoc(self.input_path)
        raise OfficeExtractionError(f"Unsupported office format: {suffix or '<none>'}")

    def _extract_doc_via_libreoffice(self) -> CleanedDocument:
        """Legacy .doc binary files are routed through a LibreOffice headless conversion
        to .docx rather than Pandoc, which handles the legacy binary format unreliably."""
        soffice = shutil.which("soffice") or shutil.which("soffice.exe") or shutil.which("libreoffice")
        if not soffice:
            raise OfficeExtractionError(
                "Converting legacy .doc files requires LibreOffice (the 'soffice' executable) "
                "to be installed and available on PATH."
            )
        with tempfile.TemporaryDirectory() as tmp_dir:
            result = subprocess.run(
                [soffice, "--headless", "--convert-to", "docx", "--outdir", tmp_dir, str(self.input_path)],
                capture_output=True,
                text=True,
                timeout=120,
            )
            converted = Path(tmp_dir) / f"{self.input_path.stem}.docx"
            if result.returncode != 0 or not converted.exists():
                raise OfficeExtractionError(
                    f"LibreOffice failed to convert {self.input_path.name} to DOCX: {result.stderr.strip()}"
                )
            original_path = self.input_path
            try:
                self.input_path = converted
                return self._extract_docx()
            finally:
                self.input_path = original_path

    def _extract_docx(self) -> CleanedDocument:
        try:
            document = Document(str(self.input_path))
        except Exception as error:
            raise OfficeExtractionError(f"Unable to read DOCX: {self.input_path}") from error

        chapters: list[ChapterContent] = []
        current_title = "Front Matter"
        current_items: list[FlowTextUnit | FlowImageUnit | FlowTableUnit] = []
        image_index = 0
        seen_relationships: set[str] = set()
        position = 0

        for block in self._iter_block_items(document):
            position += 1
            if isinstance(block, Table):
                table_unit = self._table_to_flow_unit(block, position)
                if table_unit is not None:
                    current_items.append(table_unit)
                continue

            paragraph = block
            text = self._fix_text(paragraph.text)
            level = self._heading_level(paragraph.style.name if paragraph.style else "")
            images = self._paragraph_images(paragraph, document, seen_relationships, image_index, position)
            image_index += len(images)
            if level:
                heading = FlowTextUnit(
                    kind="heading",
                    text=text or f"Section {position + 1}",
                    y=float(position),
                    page_number=1,
                    font_size=self._font_size(paragraph, level),
                    bold_ratio=self._bold_ratio(paragraph),
                    level=level,
                    is_chapter_heading=level == 1,
                )
                if level == 1 and current_items:
                    chapters.append(ChapterContent(title=current_title, items=tuple(current_items)))
                    current_items = []
                if level == 1:
                    current_title = heading.text
                current_items.append(heading)
            elif text:
                current_items.append(
                    FlowTextUnit(
                        kind="paragraph",
                        text=text,
                        y=float(position),
                        page_number=1,
                        font_size=self._font_size(paragraph, 0),
                        bold_ratio=self._bold_ratio(paragraph),
                    )
                )
            current_items.extend(images)

        if current_items or not chapters:
            chapters.append(ChapterContent(title=current_title, items=tuple(current_items)))
        return CleanedDocument(chapters=tuple(chapters))

    @staticmethod
    def _iter_block_items(parent: Any) -> Iterable[Paragraph | Table]:
        """Yield paragraphs and tables in document order (the standard python-docx recipe),
        so table content is no longer silently dropped by a paragraphs-only traversal."""
        if hasattr(parent, "element"):
            parent_element = parent.element.body
        elif isinstance(parent, _Cell):
            parent_element = parent._tc
        else:
            raise ValueError("Unsupported parent for block iteration")
        for child in parent_element.iterchildren():
            if child.tag == qn("w:p"):
                yield Paragraph(child, parent)
            elif child.tag == qn("w:tbl"):
                yield Table(child, parent)

    def _table_to_flow_unit(self, table: Table, position: int) -> FlowTableUnit | None:
        rows: list[tuple[str, ...]] = []
        for row in table.rows:
            cells: list[str] = []
            for cell in row.cells:
                cell_text_parts: list[str] = []
                for item in self._iter_block_items(cell):
                    if isinstance(item, Table):
                        nested = self._table_to_flow_unit(item, position)
                        if nested is not None:
                            cell_text_parts.append(
                                "\n".join(" | ".join(nested_row) for nested_row in nested.rows)
                            )
                    else:
                        fixed = self._fix_text(item.text)
                        if fixed:
                            cell_text_parts.append(fixed)
                cells.append("\n".join(cell_text_parts))
            rows.append(tuple(cells))
        if not rows or not any(any(cell for cell in row) for row in rows):
            return None
        return FlowTableUnit(kind="table", rows=tuple(rows), y=float(position), page_number=1, has_header_row=True)

    def _extract_with_pandoc(self, source_path: Path) -> CleanedDocument:
        with tempfile.TemporaryDirectory() as tmp_dir:
            media_dir = Path(tmp_dir) / "media"
            try:
                html_payload = pypandoc.convert_file(
                    str(source_path),
                    to="html5",
                    extra_args=[f"--extract-media={media_dir}"],
                )
            except Exception as error:
                raise OfficeExtractionError(
                    f"Unable to convert {source_path.name}; install Pandoc or provide a DOCX input"
                ) from error
            return self._extract_from_html5(html_payload, media_dir)

    def _extract_from_html5(self, html_payload: str, media_dir: Path) -> CleanedDocument:
        root = lxml_html.fromstring(html_payload)
        body = root.find("body")
        if body is None:
            body = root

        chapters: list[ChapterContent] = []
        current_title = "Front Matter"
        current_items: list[FlowTextUnit | FlowImageUnit | FlowTableUnit] = []
        position = 0

        def start_chapter(title: str) -> None:
            nonlocal current_title, current_items
            if current_items:
                chapters.append(ChapterContent(title=current_title, items=tuple(current_items)))
                current_items = []
            current_title = title

        def emit_images(element: Any, pos: int) -> None:
            for img in element.iter("img"):
                src = img.get("src", "")
                if not src:
                    continue
                image_path = (media_dir / Path(src).name) if not Path(src).is_absolute() else Path(src)
                if not image_path.exists():
                    image_path = media_dir.parent / src
                if not image_path.exists():
                    continue
                extension = image_path.suffix.lower().lstrip(".") or "png"
                current_items.append(
                    FlowImageUnit(
                        kind="image",
                        y=float(pos) + 0.1,
                        page_number=1,
                        image_bytes=image_path.read_bytes(),
                        extension=extension,
                        alt=img.get("alt", ""),
                        full_width=True,
                    )
                )

        for element in body.iter():
            if element.tag in {"h1", "h2", "h3", "h4", "h5", "h6"}:
                position += 1
                level = min(3, int(element.tag[1]))
                text = self._fix_text(element.text_content())
                heading = FlowTextUnit(
                    kind="heading",
                    text=text or f"Section {position}",
                    y=float(position),
                    page_number=1,
                    font_size=max(12.0, 24.0 - level * 2.0),
                    bold_ratio=1.0,
                    level=level,
                    is_chapter_heading=level == 1,
                )
                if level == 1:
                    start_chapter(heading.text)
                current_items.append(heading)
            elif element.tag in {"p", "li"}:
                position += 1
                text = self._fix_text(element.text_content())
                if text:
                    prefix = "• " if element.tag == "li" else ""
                    current_items.append(
                        FlowTextUnit(
                            kind="paragraph",
                            text=f"{prefix}{text}",
                            y=float(position),
                            page_number=1,
                            font_size=12.0,
                            bold_ratio=0.0,
                        )
                    )
                emit_images(element, position)
            elif element.tag == "table":
                position += 1
                table_unit = self._html_table_to_flow_unit(element, position)
                if table_unit is not None:
                    current_items.append(table_unit)
            elif element.tag == "img" and element.getparent() is not None and element.getparent().tag not in {"p", "li"}:
                position += 1
                emit_images(element.getparent(), position)

        if current_items or not chapters:
            chapters.append(ChapterContent(title=current_title, items=tuple(current_items)))
        return CleanedDocument(chapters=tuple(chapters))

    def _html_table_to_flow_unit(self, table_element: Any, position: int) -> FlowTableUnit | None:
        rows: list[tuple[str, ...]] = []
        has_header_row = table_element.find(".//th") is not None
        for row_element in table_element.iter("tr"):
            cells = [self._fix_text(cell.text_content()) for cell in row_element.iter("th", "td")]
            if cells:
                rows.append(tuple(cells))
        if not rows or not any(any(cell for cell in row) for row in rows):
            return None
        return FlowTableUnit(kind="table", rows=tuple(rows), y=float(position), page_number=1, has_header_row=has_header_row)

    def _paragraph_images(self, paragraph: Any, document: Any, seen: set[str], start_index: int, position: int) -> list[FlowImageUnit]:
        images: list[FlowImageUnit] = []
        for run in paragraph.runs:
            for blip in run._element.xpath('.//*[local-name()="blip"]'):
                relationship = blip.get(qn("r:embed"))
                if not relationship or relationship in seen:
                    continue
                part = document.part.related_parts.get(relationship)
                if part is None or not hasattr(part, "blob"):
                    continue
                seen.add(relationship)
                extension = self._extension_for_content_type(getattr(part, "content_type", "image/png"))
                images.append(
                    FlowImageUnit(
                        kind="image",
                        y=float(position) + 0.1,
                        page_number=1,
                        image_bytes=part.blob,
                        extension=extension,
                        alt="Embedded document image",
                        full_width=True,
                    )
                )
        return images

    @staticmethod
    def _heading_level(style_name: str) -> int:
        match = re.search(r"heading\s*([1-3])", style_name or "", flags=re.IGNORECASE)
        return int(match.group(1)) if match else 0

    @staticmethod
    def _font_size(paragraph: Any, level: int) -> float:
        sizes = [run.font.size.pt for run in paragraph.runs if run.font.size is not None]
        return max(sizes, default=max(12.0, 22.0 - level * 2.0))

    @staticmethod
    def _bold_ratio(paragraph: Any) -> float:
        runs = [run for run in paragraph.runs if run.text.strip()]
        return sum(bool(run.bold) for run in runs) / len(runs) if runs else 0.0

    @staticmethod
    def _fix_text(text: str) -> str:
        return " ".join(ftfy.fix_text(text, uncurl_quotes=False).split())

    @staticmethod
    def _extension_for_content_type(content_type: str) -> str:
        return {
            "image/jpeg": "jpg",
            "image/gif": "gif",
            "image/webp": "webp",
            "image/svg+xml": "svg",
        }.get(content_type, "png")
